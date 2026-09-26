#!/usr/bin/env python3
"""Extrae de un PPT la estructura que necesita construir_fichas.py (se ejecuta en el Mac, con python-pptx).

Uso: python3 extraer_ppt.py "TEMA 3. Creta y Grecia.pptx" t3_slides.json [carpeta_imagenes]
Si se indica carpeta_imagenes, guarda además cada imagen incrustada como sNN_K.ext para elegir cuáles van a la app.
"""
import sys, os, json
from pptx import Presentation
from pptx.util import Emu

src, out = sys.argv[1], sys.argv[2]
imgdir = sys.argv[3] if len(sys.argv) > 3 else None
if imgdir: os.makedirs(imgdir, exist_ok=True)
prs = Presentation(src)

def shapes_of(s):
    res = []
    def walk(shs):
        for sh in shs:
            if sh.shape_type == 6: walk(sh.shapes)
            else: res.append(sh)
    walk(s.shapes); return res

data = []
for i, s in enumerate(prs.slides, 1):
    title = None; paras = []; n_img = 0
    shs = shapes_of(s)
    shs.sort(key=lambda sh: ((sh.top or 0) // Emu(400000), (sh.left or 0)))
    for sh in shs:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if not t: continue
            is_title = sh.is_placeholder and sh.placeholder_format.type is not None and \
                str(sh.placeholder_format.type).split(".")[-1].split(" ")[0] in ("TITLE", "CENTER_TITLE")
            if is_title and title is None: title = t
            else:
                for p in sh.text_frame.paragraphs:
                    pt = "".join(r.text for r in p.runs).strip()
                    if pt: paras.append(pt)
        if sh.shape_type == 13:
            n_img += 1
            if imgdir:
                try:
                    im = sh.image
                    with open(os.path.join(imgdir, f"s{i:02d}_{n_img}.{im.ext}"), "wb") as f: f.write(im.blob)
                except Exception:
                    pass
    notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
    data.append({"slide": i, "title": title, "paras": paras, "n_img": n_img, "notes": notes})
json.dump(data, open(out, "w"), ensure_ascii=False, indent=1)
print(len(data), "diapositivas;", sum(1 for d in data if d["title"]), "con título")
